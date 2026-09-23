from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import PurePosixPath
from defusedxml import ElementTree as DefusedET
from sqlalchemy.orm import Session

from .specs import parse_document_spec, semantic_anchor_text, spreadsheet_rows
from .renderers.docx import render_docx
from .renderers.pdf import PdfRendererUnavailable, render_pdf
from .renderers.pptx import PptxRendererUnavailable, render_pptx
from .renderers.xlsx import XlsxRendererUnavailable, render_xlsx
from .storage import BlobStorage, BlobStorageError, build_blob_storage
from ..models import ArtifactAsset, ArtifactJob, ArtifactUsageEvent

DOCX_MIME="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME="application/pdf"
PPTX_MIME="application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_MIME="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

FORMAT_MAP={
    "docx":(".docx",DOCX_MIME),
    "pdf":(".pdf",PDF_MIME),
    "pptx":(".pptx",PPTX_MIME),
    "xlsx":(".xlsx",XLSX_MIME),
}

class ArtifactError(RuntimeError): pass
class ArtifactFormatUnsupported(ArtifactError): pass
class ArtifactValidationFailed(ArtifactError): pass
class ArtifactStorageError(ArtifactError): pass

@dataclass(frozen=True, slots=True)
class ValidatedArtifact:
    format: str
    filename: str
    mime_type: str
    data: bytes
    sha256: str
    semantic_text: str
    renderer: str

def _safe_filename(name:str, ext:str)->str:
    base=PurePosixPath((name or "").replace("\\","/")).name.strip()
    if not base or base in {".",".."}: base="glip-artifact"
    base=re.sub(r"[^A-Za-z0-9._ -]+","_",base).strip(" ._") or "glip-artifact"
    stem=base.rsplit(".",1)[0] if "." in base else base
    return f"{stem[:120]}{ext}"

def _extract_docx_text(data:bytes)->str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if z.testzip() is not None:
                raise ArtifactValidationFailed("ARTIFACT_DOCX_CRC_FAILED")
            required={"[Content_Types].xml","_rels/.rels","word/document.xml"}
            if not required.issubset(set(z.namelist())):
                raise ArtifactValidationFailed("ARTIFACT_DOCX_STRUCTURE_INVALID")
            raw=z.read("word/document.xml")
        root=DefusedET.fromstring(raw,forbid_dtd=True,forbid_entities=True,forbid_external=True)
    except ArtifactValidationFailed:
        raise
    except Exception as exc:
        raise ArtifactValidationFailed("ARTIFACT_DOCX_VALIDATION_FAILED") from exc
    ns="{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    text="\n".join((node.text or "") for node in root.iter(f"{ns}t")).strip()
    if not text: raise ArtifactValidationFailed("ARTIFACT_DOCX_SEMANTIC_EMPTY")
    return text

def _extract_pdf_text(data:bytes)->str:
    try:
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data),strict=True)
        text="\n".join(page.extract_text() or "" for page in reader.pages).strip()
    except Exception as exc:
        raise ArtifactValidationFailed("ARTIFACT_PDF_VALIDATION_FAILED") from exc
    if not text: raise ArtifactValidationFailed("ARTIFACT_PDF_SEMANTIC_EMPTY")
    return text

def _extract_pptx_text(data:bytes)->str:
    try:
        from pptx import Presentation
        presentation=Presentation(io.BytesIO(data))
        values=[]
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape,"has_text_frame",False):
                    values.append(shape.text or "")
        text="\n".join(values).strip()
    except Exception as exc:
        raise ArtifactValidationFailed("ARTIFACT_PPTX_VALIDATION_FAILED") from exc
    if not text: raise ArtifactValidationFailed("ARTIFACT_PPTX_SEMANTIC_EMPTY")
    return text

def _extract_xlsx_text(data:bytes)->str:
    try:
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        values=[]
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells=[str(value) for value in row if value is not None and str(value).strip()]
                if cells: values.append(" | ".join(cells))
        wb.close()
        text="\n".join(values).strip()
    except Exception as exc:
        raise ArtifactValidationFailed("ARTIFACT_XLSX_VALIDATION_FAILED") from exc
    if not text: raise ArtifactValidationFailed("ARTIFACT_XLSX_SEMANTIC_EMPTY")
    return text

def render_and_validate(*,format:str,content:str,filename:str|None=None)->ValidatedArtifact:
    fmt=(format or "").strip().lower()
    if fmt not in FORMAT_MAP: raise ArtifactFormatUnsupported("ARTIFACT_FORMAT_UNSUPPORTED")
    normalized=(content or "").strip()
    if not normalized: raise ArtifactValidationFailed("ARTIFACT_EMPTY_CONTENT")
    ext,mime=FORMAT_MAP[fmt]
    spec=parse_document_spec(normalized)
    if not spec.semantic_text.strip(): raise ArtifactValidationFailed("ARTIFACT_SEMANTIC_EMPTY")
    try:
        if fmt=="docx":
            data=render_docx(spec); observed=_extract_docx_text(data); renderer="glip_docx_v1"
        elif fmt=="pdf":
            data=render_pdf(spec); observed=_extract_pdf_text(data); renderer="glip_pdf_v1"
        elif fmt=="pptx":
            data=render_pptx(spec); observed=_extract_pptx_text(data); renderer="glip_pptx_v1"
        else:
            data=render_xlsx(normalized,spec); observed=_extract_xlsx_text(data); renderer="glip_xlsx_v1"
    except (PdfRendererUnavailable,PptxRendererUnavailable,XlsxRendererUnavailable) as exc:
        raise ArtifactFormatUnsupported(str(exc)) from exc
    probe=(("\n".join(" | ".join(r) for r in spreadsheet_rows(normalized,spec)[:3]))
           if fmt=="xlsx" else semantic_anchor_text(spec,limit=120))
    probe=re.sub(r"\s+"," ",probe).strip()
    normalized_observed=re.sub(r"\s+"," ",observed).strip()
    if probe and probe not in normalized_observed:
        anchors=[x.strip() for x in re.split(r"\n+|\s*\|\s*",spec.semantic_text) if x.strip()][:3]
        if anchors and not all(re.sub(r"\s+"," ",a)[:80] in normalized_observed for a in anchors):
            raise ArtifactValidationFailed("ARTIFACT_SEMANTIC_MISMATCH")
    return ValidatedArtifact(
        format=fmt,filename=_safe_filename(filename or f"glip-{fmt}",ext),
        mime_type=mime,data=data,sha256=hashlib.sha256(data).hexdigest(),
        semantic_text=observed,renderer=renderer,
    )

def persist_artifact(
    db:Session,*,settings,tenant_id:str,project_id:str,actor_id:str,
    validated:ValidatedArtifact,classification:str="project_internal",
    source_kind:str="user_content",source_ref:str|None=None,
    storage:BlobStorage|None=None,
)->tuple[ArtifactJob,ArtifactAsset]:
    if not getattr(settings,"artifacts_enabled",False):
        raise ArtifactError("ARTIFACTS_DISABLED")
    storage=storage or build_blob_storage(settings)
    job=ArtifactJob(
        tenant_id=tenant_id,project_id=project_id,artifact_kind="document",
        requested_format=validated.format,status="rendered",source_kind=source_kind,
        source_ref=source_ref,request_json={"filename":validated.filename},
        billing_scope="glip",provider_profile_ref=getattr(settings,"artifact_provider_profile_ref","glip-local"),
        requested_by=actor_id,
    )
    db.add(job); db.flush()
    artifact_id=str(uuid.uuid4())
    key=f"{tenant_id}/{project_id}/artifacts/{artifact_id}-{validated.filename}"
    try:
        created=storage.put_if_absent(key,validated.data,content_type=validated.mime_type)
        if not created: raise ArtifactStorageError("ARTIFACT_STORAGE_KEY_COLLISION")
        stored=storage.get(key)
        if hashlib.sha256(stored).hexdigest()!=validated.sha256:
            raise ArtifactValidationFailed("ARTIFACT_STORED_HASH_MISMATCH")
    except (BlobStorageError,ArtifactError) as exc:
        job.status="failed"; job.error_code=str(exc)
        raise
    asset=ArtifactAsset(
        id=artifact_id,tenant_id=tenant_id,project_id=project_id,job_id=job.id,
        artifact_kind="document",format=validated.format,filename=validated.filename,
        mime_type=validated.mime_type,storage_key=key,sha256=validated.sha256,
        size_bytes=len(validated.data),classification=classification,
        metadata_json={"renderer":validated.renderer,"validated":True,"billing_scope":"glip"},
        version=1,created_by=actor_id,
    )
    db.add(asset)
    db.add(ArtifactUsageEvent(
        tenant_id=tenant_id,project_id=project_id,artifact_job_id=job.id,
        billing_scope="glip",usage_type="document_render",provider="local_renderer",
        provider_profile_ref=getattr(settings,"artifact_provider_profile_ref","glip-local"),
        units_json={"format":validated.format,"bytes":len(validated.data)},
        cost_microusd=None,
    ))
    job.status="completed"; job.completed_at=datetime.now(timezone.utc)
    db.flush()
    return job,asset

def artifact_capability_manifest(settings)->dict:
    enabled=bool(getattr(settings,"artifacts_enabled",False))
    return {
        "schema_version":"glip.artifact-capabilities.v1",
        "enabled":enabled,
        "billing_scope":"glip",
        "provider_profile_ref":getattr(settings,"artifact_provider_profile_ref","glip-local"),
        "formats":{
            "docx":{"supported":True,"enabled":enabled,"renderer":"glip_docx_v1"},
            "pdf":{"supported":True,"enabled":enabled,"renderer":"glip_pdf_v1"},
            "pptx":{"supported":True,"enabled":enabled,"renderer":"glip_pptx_v1"},
            "xlsx":{"supported":True,"enabled":enabled,"renderer":"glip_xlsx_v1"},
            "image":{"supported":False,"enabled":False,"status":"IMAGE_PROVIDER_PENDING"},
        },
        "architecture":{
            "ifc":{"first_class_source":True,"semantic_parser":"PENDING_GATE_2"},
            "dwg":{"first_class_source":True,"automation":"PENDING_GATE_2"},
            "dxf":{"first_class_source":True,"parser":"PENDING_GATE_2"},
            "glb_gltf":{"distribution_format":True},
            "render_worker":{"engine":"blender","status":"PENDING_GATE_4"},
        },
    }
